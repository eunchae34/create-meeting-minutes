#!/usr/bin/env python3
"""
회의록 양식.docx 템플릿에 구조화된 회의 데이터(JSON)를 채워 넣어
표준 형식의 회의록 .docx 파일을 생성한다.

사용법:
    py build_meeting_minutes.py --template <template.docx> --data <data.json> --output <output.docx>

data.json 스키마는 이 디렉토리의 sample_data.json 참고.
"""
import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.shared import OxmlElement
from docx.shared import Pt, Twips

STYLE_BODY = "표내용"    # 값 셀 내부 텍스트

# .../회의록/.claude/skills/create-meeting-minutes/scripts/build_meeting_minutes.py
# -> parents[4] == .../회의록  (이 스킬이 속한 프로젝트 폴더, 실제 회의록 파일들이 모여 있는 곳)
PROJECT_ROOT = Path(__file__).resolve().parents[4]

HEADING_SIZE = 14
SUBHEADING_SIZE = 11
BODY_SIZE = 9

# 사용자 요청으로 서체를 맑은 고딕으로 고정한다. 템플릿 스타일(표내용/표목차/Normal)마다
# 기본 서체가 제각각이라 그대로 두면 문서 안에서 서체가 섞여 보이는 문제가 있었음.
FONT_NAME = "맑은 고딕"


def set_font(run, name=FONT_NAME):
    """run의 서체를 라틴/이스트에이시안 둘 다 지정한다.
    run.font.name만 설정하면 한글 부분(w:eastAsia)에는 적용되지 않아 Word에서
    스타일 기본 서체로 되돌아가는 경우가 있어, rFonts를 직접 만져 양쪽 다 고정한다."""
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


# --- Word 네이티브 다단계 번호 매기기 ---------------------------------------
#
# 사용자가 실제 회의록 파일의
# numbering.xml을 직접 확인해 요청한 형태: bullet 항목이 수동으로 탭+텍스트를 붙인 게
# 아니라 진짜 Word 목록(w:numPr)이었다 - 그래서 "I. II. III."는 왼쪽 정렬, "i. ii. iii."는
# 오른쪽 정렬이고, 번호 폭이 달라져도(예: "IV." vs "V.") 본문 시작 위치가 흔들리지 않는다.
#
# 처음엔 abstractNum 하나를 여러 목록이 공유하고 numId만 새로 발급했는데, 그러면 Word가
# 안건이 바뀌어도 번호를 이어서 매겨버렸다(예: 2번 안건의 첫 항목이 I이 아니라 IV로 시작).
# 실제 회의록 파일의 numbering.xml을 다시 보니 목록마다 abstractNum 자체가 전부 별도로
# 복제되어 있었다(같은 abstractNumId를 공유하는 경우가 없음) - Word가 새 목록을 만들 때
# 하는 방식 그대로다. 그래서 새 목록이 시작될 때마다 abstractNum부터 통째로 새로 만든다.
#
# 회의록 양식 템플릿에는 이미 To do list용 체크박스 목록(numId=1, Wingdings 문자)이
# 정의되어 있어 그대로 재사용한다.

TODO_CHECKBOX_NUM_ID = 1  # 템플릿에 이미 정의되어 있는 체크박스 목록의 numId


def _numbering_root(doc):
    return doc.part.numbering_part.element


def _next_abstract_num_id(numbering_root):
    ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering_root.findall(qn("w:abstractNum"))]
    return (max(ids) + 1) if ids else 0


def _add_abstract_num(numbering_root, abstract_id, num_fmt, lvl_jc, ind_left, ind_hanging):
    """실제 회의록 파일의 abstractNum 정의를 그대로 복제해 새로 추가한다."""
    xml = (
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0">'
        f'<w:start w:val="1"/>'
        f'<w:numFmt w:val="{num_fmt}"/>'
        f'<w:lvlText w:val="%1."/>'
        f'<w:lvlJc w:val="{lvl_jc}"/>'
        f'<w:pPr><w:ind w:left="{ind_left}" w:hanging="{ind_hanging}"/></w:pPr>'
        f"</w:lvl>"
        f"</w:abstractNum>"
    )
    el = parse_xml(xml)
    # abstractNum은 반드시 모든 num보다 앞에 와야 하므로 맨 앞에 삽입한다.
    numbering_root.insert(0, el)
    return el


def new_upper_roman_list(doc):
    """대문자 로마 숫자("I. II. III. ...") 목록을 abstractNum부터 새로 만들어 numId를 반환한다.
    매번 abstractNum까지 새로 만들어야 항상 I부터 다시 시작된다 (위 설명 참고)."""
    numbering_root = _numbering_root(doc)
    abstract_id = _next_abstract_num_id(numbering_root)
    _add_abstract_num(numbering_root, abstract_id, "upperRoman", "left", 880, 440)
    return numbering_root.add_num(abstract_id).numId


def new_lower_roman_list(doc):
    """소문자 로마 숫자("i. ii. iii. ...") 하위 목록을 abstractNum부터 새로 만들어 numId를 반환한다."""
    numbering_root = _numbering_root(doc)
    abstract_id = _next_abstract_num_id(numbering_root)
    _add_abstract_num(numbering_root, abstract_id, "lowerRoman", "right", 1320, 440)
    return numbering_root.add_num(abstract_id).numId


def set_numbering(paragraph, num_id, ilvl=0):
    """문단에 실제 Word 목록 번호(numPr)를 적용한다."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


def clear_cell(cell):
    """셀의 모든 문단을 지우고, 런(run)이 비워진 첫 번째 문단만 남겨 반환한다.
    첫 번째 문단 객체 자체는 유지하므로 원래 스타일(표내용/표목차/Normal 등)이 그대로 보존된다."""
    paragraphs = cell.paragraphs
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    return first


def add_run(paragraph, text, bold=False, size_pt=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    set_font(run)
    return run


def set_plain_field(cell, text, bold=False, size_pt=None):
    """회의 일시 / 부서 / 작성자 / 회의 장소 처럼 한 줄짜리 값을 채운다."""
    p = clear_cell(cell)
    add_run(p, text, bold=bold, size_pt=size_pt)


def set_multiline_field(cell, lines, style_name=STYLE_BODY, size_pt=BODY_SIZE):
    """참석자 / 회의 주제처럼 줄 단위로 여러 문단이 들어갈 수 있는 값을 채운다."""
    if isinstance(lines, str):
        lines = [lines]
    first = clear_cell(cell)
    if not lines:
        return
    add_run(first, lines[0], bold=False, size_pt=size_pt)
    for line in lines[1:]:
        p = cell.add_paragraph(style=style_name)
        add_run(p, line, bold=False, size_pt=size_pt)


QNA_ANSWER_INDENT_TWIPS = 880  # upperRoman 목록의 ind left와 맞춰서, Q 문단의 본문 시작 위치에 A를 정렬


def build_content_cell(doc, cell, blocks):
    """회의 내용 셀을 content_blocks 리스트로 채운다.

    block type:
      heading    - 대분류 제목 (예: "📄 주요 안건", "❔Q&A") -> 14pt bold.
                   이모지가 필요하면 text 안에 직접 포함시킨다 (스크립트가 자동으로 붙여주지 않음).
      subheading - 번호 매긴 안건 제목 (예: "1. 서비스 모델 논의") -> 11pt bold
      bullet     - 세부 내용 한 줄. {"type": "bullet", "text": "...", "sub_bullets": [...]}
                   Word 네이티브 목록(w:numPr)으로 대문자 로마 숫자("I. II. III. ...")가 자동으로
                   매겨진다 (heading/subheading을 새로 만나면 새 목록으로 다시 I부터 시작).
                   "sub_bullets"(선택)를 주면 그 bullet 밑에 별도의 네이티브 목록으로 소문자
                   로마 숫자("i. ii. ...")가 매겨진 하위 목록이 생긴다 (bullet마다 새로 i부터 시작).
      qna        - {"q": "...", "q_speaker": "최지훈 리더", "a": "...", "a_speaker": "이영수 팀장"}
                   Q는 bullet과 같은 대문자 로마 숫자 목록에 "Qn. <질문> (<q_speaker>)"로 들어가고
                   (heading/subheading을 새로 만나기 전까지는 같은 목록이 계속 이어진다 - 즉 같은
                   Q&A 구간 안의 질문들은 I, II, III...로 쭉 이어지고 새 구간에서만 다시 I부터 시작),
                   A는 번호 없이 "A. <답변> (<a_speaker>)"로 Q의 본문 시작 위치에 맞춰 들여쓴다.
                   q_speaker/a_speaker는 선택 — 주면 문장 끝에 "(이름)"으로 붙는다.
      blank      - 빈 줄 (섹션 구분용)
    """
    first = clear_cell(cell)
    first_used = False
    current_upper_num_id = None  # 현재 subheading/heading 구간에서 쓰고 있는 목록의 numId
    qna_no = 0  # 현재 Q&A 구간에서 몇 번째 질문인지 ("Qn."에 쓰임)

    def next_paragraph():
        nonlocal first_used
        if not first_used:
            first_used = True
            return first
        return cell.add_paragraph(style=STYLE_BODY)

    for block in blocks:
        btype = block.get("type")
        if btype == "heading":
            current_upper_num_id = None
            qna_no = 0
            p = next_paragraph()
            add_run(p, block["text"], bold=True, size_pt=HEADING_SIZE)
        elif btype == "subheading":
            current_upper_num_id = None
            qna_no = 0
            p = next_paragraph()
            add_run(p, block["text"], bold=True, size_pt=SUBHEADING_SIZE)
        elif btype == "bullet":
            if current_upper_num_id is None:
                current_upper_num_id = new_upper_roman_list(doc)
            p = next_paragraph()
            set_numbering(p, current_upper_num_id)
            add_run(p, block["text"], bold=False, size_pt=BODY_SIZE)
            sub_bullets = block.get("sub_bullets", [])
            if sub_bullets:
                sub_num_id = new_lower_roman_list(doc)
                for sub_text in sub_bullets:
                    sp = next_paragraph()
                    set_numbering(sp, sub_num_id)
                    add_run(sp, sub_text, bold=False, size_pt=BODY_SIZE)
        elif btype == "qna":
            if current_upper_num_id is None:
                current_upper_num_id = new_upper_roman_list(doc)
            qna_no += 1
            q_speaker = block.get("q_speaker", "")
            q_suffix = f" ({q_speaker})" if q_speaker else ""
            qp = next_paragraph()
            set_numbering(qp, current_upper_num_id)
            add_run(qp, f"Q{qna_no}. {block['q']}{q_suffix}", bold=False, size_pt=BODY_SIZE)

            a_speaker = block.get("a_speaker", "")
            a_suffix = f" ({a_speaker})" if a_speaker else ""
            ap = next_paragraph()
            ap.paragraph_format.left_indent = Twips(QNA_ANSWER_INDENT_TWIPS)
            add_run(ap, f"A. {block['a']}{a_suffix}", bold=False, size_pt=BODY_SIZE)
        elif btype == "blank":
            next_paragraph()  # 빈 문단만 추가
        else:
            raise ValueError(f"알 수 없는 block type: {btype}")


def build_todo_cell(cell, todos):
    """기타 셀. 템플릿에 이미 있는 '■ To do list' 헤더는 그대로 두고,
    그 아래에 항목별로 체크박스 목록(템플릿에 이미 정의된 numId=1) 문단을 추가한다."""
    paragraphs = cell.paragraphs
    # 첫 번째 문단('■ To do list')만 남기고 나머지(빈 문단들) 제거
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for item in todos or []:
        p = cell.add_paragraph(style=STYLE_BODY)
        set_numbering(p, TODO_CHECKBOX_NUM_ID)
        add_run(p, item, bold=False, size_pt=BODY_SIZE)


def parse_yymmdd(date_str):
    """'2026년 8월 14일' 같은 문자열에서 YYMMDD를 뽑아낸다. 실패하면 None."""
    m = re.search(r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일", date_str)
    if not m:
        return None
    yy, mm, dd = m.group(1)[-2:], int(m.group(2)), int(m.group(3))
    return f"{yy}{mm:02d}{dd:02d}"


def build(template_path, data_path, output_path=None):
    with open(data_path, encoding="utf-8") as f:
        data = json.load(f)

    doc = Document(template_path)
    t0 = doc.tables[0]  # 회의 일시 / 장소 / 참석자 / 부서 / 작성자
    t1 = doc.tables[1]  # 회의 주제 / 회의 내용 / 기타

    # --- table0: 회의 일시 / 부서 / 작성자 (row0), 회의 장소 (row1), 참석자 (row2) ---
    set_plain_field(t0.rows[0].cells[1], data["date"])
    set_plain_field(t0.rows[0].cells[3], data["department"])
    set_plain_field(t0.rows[0].cells[5], data["author"])
    set_plain_field(t0.rows[1].cells[1], data["place"], bold=True, size_pt=8.5)
    set_multiline_field(t0.rows[2].cells[1], data["attendees"])

    # --- table1: 회의 주제 / 회의 내용 / 기타(To do list) ---
    set_multiline_field(t1.rows[0].cells[1], data["topic"])
    build_content_cell(doc, t1.rows[1].cells[1], data["content_blocks"])
    build_todo_cell(t1.rows[2].cells[1], data.get("todos", []))

    if not output_path:
        output_path = data.get("output_filename")
    if not output_path:
        yymmdd = parse_yymmdd(data["date"]) or "yymmdd"
        topic = data.get("topic", "회의록").strip()
        output_path = f"{yymmdd}_{topic} 회의록.docx"

    out = Path(output_path)
    if not out.is_absolute():
        # 파일명만 주어졌거나 상대 경로면, 실행 위치와 무관하게 항상 이 프로젝트의
        # 회의록 폴더(PROJECT_ROOT) 기준으로 저장한다 - 기존 회의록 파일들과 같은 곳에 쌓이도록.
        out = PROJECT_ROOT / out
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return str(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--output", default=None)
    args = ap.parse_args()

    saved = build(args.template, args.data, args.output)
    print(f"OK: {saved}")


if __name__ == "__main__":
    main()
